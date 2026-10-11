# -*- coding: utf-8 -*-
"""
Point d'entree : genere le deck Africa CDC (pptx) d'un SitRep RDC, verifie chaque chiffre, ecrit les rapports.

  python3 scripts/python/deck/build_deck.py --sitrep 147 \\
      --pdf-dir data/pdf --template data/africa_cdc_deck/templates/deck_template.pptx \\
      --state data/africa_cdc_deck/state.json --out-dir outputs/deck_private [--narrative auto|off] [--commit-state]

Codes retour : 0 deck genere et verifie ; 2 SitRep illisible / precedent indisponible (aucun deck) ;
               3 erreur de gabarit ; 4 verification des chiffres echouee (deck produit mais NON valide -> ne pas envoyer).
Fichiers ecrits dans --out-dir : deck pptx, Controle_<...>.csv, deck_report.json.
"""
import argparse
import csv
import datetime as dt
import glob
import json
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import sitrep_parse as sp  # noqa: E402
import deck_build as db  # noqa: E402

MOIS_FR = ["", "Janvier", "Fevrier", "Mars", "Avril", "Mai", "Juin", "Juillet", "Aout", "Septembre", "Octobre", "Novembre", "Decembre"]

DEFAULT_CONFIG = {
    # valeurs RETENUES des publications precedentes (absentes des SitRep RDC) ; a mettre a jour a la main si Africa CDC publie du nouveau
    "uganda": {"cases": 20, "deaths": 2, "recoveries": 18},
    "hcw": {"drc_infections": 266, "uganda_infections": 4, "deaths": 49},
}


def pdftotext(path):
    r = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True)
    if r.returncode != 0:
        raise sp.ParseError("pdftotext a echoue sur %s : %s" % (path, r.stderr.decode("utf-8", "replace")[:200]))
    return r.stdout.decode("utf-8", "replace")


def find_pdf(pdf_dir, n):
    for pat in ("SitRep_%d_2026.pdf", "SitRep_%03d_2026.pdf", "PREIS_DRC_Ebola_SitRep_%d.pdf", "PREIS_DRC_Ebola_SitRep_%03d.pdf"):
        fp = os.path.join(pdf_dir, pat % n)
        if os.path.isfile(fp):
            return fp
    c = sorted(glob.glob(os.path.join(pdf_dir, "*%d*.pdf" % n)))
    return c[0] if c else None


def load_history(pdf_dir):
    """cumuls (cas, deces) par date de rapport, lus de TOUS les PDF disponibles (officiels)."""
    h, src = {}, {}
    for fp in sorted(glob.glob(os.path.join(pdf_dir, "*.pdf"))):
        try:
            k = sp.quick_kpi(pdftotext(fp))
        except Exception:
            continue
        if not k:
            continue
        no, d, c, de = k
        key = d.isoformat()
        if key in h and h[key] != (c, de):
            # deux PDF pour la meme date avec des cumuls differents : on garde le plus recent (numero le plus grand)
            if no < src[key]:
                continue
        h[key], src[key] = (c, de), no
    return h


def find_prev(pdf_dir, D):
    """SitRep dont la date de rapport est la veille : analyse STRICTE (meme exigence que le courant)."""
    want = D - dt.timedelta(days=1)
    for fp in sorted(glob.glob(os.path.join(pdf_dir, "*.pdf"))):
        try:
            t = pdftotext(fp)
            k = sp.quick_kpi(t)
        except Exception:
            continue
        if k and k[1] == want:
            S = sp.parse_sitrep(t)
            return S, fp
    return None, None


def load_json(fp, default):
    try:
        return json.load(open(fp, encoding="utf-8"))
    except Exception:
        return default


def verify_deck(path, S, P, cx_week):
    """Relit le pptx ENREGISTRE (sans passer par le code de construction) et compare les valeurs cles au SitRep."""
    from pptx import Presentation
    p = Presentation(path)
    fails, ok = [], 0

    def txt(sn, sid, para=None):
        for sh in p.slides[sn - 1].shapes:
            if sh.shape_id == sid and sh.has_text_frame:
                return sh.text_frame.text if para is None else sh.text_frame.paragraphs[para].text
        return None

    def check(label, got, want):
        nonlocal ok
        if got == want:
            ok += 1
        else:
            fails.append("%s : deck=%r attendu=%r" % (label, got, want))

    def has(label, sn, sid, frag, para=None):
        nonlocal ok
        t = txt(sn, sid, para)
        if t is not None and frag in t:
            ok += 1
        else:
            fails.append("%s : %r absent de %r" % (label, frag, t))

    n0 = db.n0
    D = S["report_date"]
    c, d = S["cases_cum"], S["deaths_cum"]
    check("couverture : numero d'Issue", txt(1, 4), "SITUATION REPORT NO. %d" % (S["sitrep_no"] - 10))
    check("couverture : date", txt(1, 9), "%02d %s %d" % (D.day, db.MONTHS[D.month], D.year))
    for sn, ids in ((5, (9, 10, 13, 14)), (6, (24, 25, 29, 30))):
        check("diapo %d nouveaux cas" % sn, txt(sn, ids[0]), str(S["new_cases"]))
        check("diapo %d cumul cas" % sn, txt(sn, ids[1]), "Cumulative %s (DRC)" % n0(c))
        check("diapo %d deces" % sn, txt(sn, ids[2]), str(S["new_deaths"]))
        check("diapo %d cumul deces" % sn, txt(sn, ids[3]), "Cumulative %s" % n0(d))
        check("diapo %d deces communautaires" % sn, txt(sn, 20), str(S["new_comm_deaths"]))
    # tableau diapo 5
    t = None
    for sh in p.slides[4].shapes:
        if sh.has_table and sh.shape_id == 2:
            t = sh.table
    dbp = db._deaths_by_prov(S)
    if t is not None and dbp is not None:
        for i in range(1, len(t.rows)):
            nm = t.cell(i, 0).text
            if nm == "Total":
                check("diapo 5 total", (t.cell(i, 1).text, t.cell(i, 2).text), (str(S["new_cases"]), str(S["new_deaths"])))
            else:
                pn = sp.prov_name(nm)
                if pn in S["t1"]:
                    check("diapo 5 %s" % nm, (t.cell(i, 1).text, t.cell(i, 2).text), (str(S["t1"][pn]["new"]), str(dbp.get(pn, 0))))
    # diapo 8
    check("diapo 8 nouveaux cas", txt(8, 10), str(S["new_cases"]))
    check("diapo 8 cumul cas", txt(8, 35), n0(c))
    check("diapo 8 cumul deces", txt(8, 39), n0(d))
    check("diapo 8 isolement", txt(8, 43), n0(S["isolation_total"]))
    ct = S.get("contacts")
    if ct:
        check("diapo 8 suivi", txt(8, 47), "%.1f%%" % ct["pct"])
        check("diapo 8 contacts", txt(8, 48), "%s of %s" % (n0(ct["seen"]), n0(ct["listed"])))
    # diapo 9 : tableau des provinces
    for sh in p.slides[8].shapes:
        if sh.has_table and sh.shape_id == 148:
            tb = sh.table
            for i in range(1, len(tb.rows)):
                nm = tb.cell(i, 0).text
                if nm == "Total":
                    want = ("%d/%d" % (S["hz_affected"], S["hz_total"]), str(S["new_cases"]), n0(c), "%s (%.1f%%)" % (n0(d), S["nat_cfr_calc"]))
                else:
                    pn = sp.prov_name(nm.replace("DRC · ", ""))
                    if pn not in S["t1"]:
                        continue
                    v = S["t1"][pn]
                    want = ("%d/%d" % (v["hz_aff"], v["hz_tot"]), str(v["new"]), n0(v["cases"]), "%s (%.1f%%)" % (n0(v["deaths"]), v["cfr"]))
                check("diapo 9 %s" % nm, tuple(tb.cell(i, j).text for j in range(1, 5)), want)
    # tuiles des diapos 11-13 (ids ...) et 14-15
    for sn, pn in ((11, "Ituri"), (12, "Haut-Uélé"), (13, "Tshopo")):
        v = S["t1"][pn]
        check("diapo %d cas" % sn, txt(sn, 9), n0(v["cases"]))
        check("diapo %d deces" % sn, txt(sn, 12), n0(v["deaths"]))
        check("diapo %d letalite" % sn, txt(sn, 14), "%.1f%%" % v["cfr"])
        check("diapo %d zones" % sn, txt(sn, 16), "%d / %d" % (v["hz_aff"], v["hz_tot"]))
    for sn, pn in ((14, "Bas-Uélé"), (15, "Sud Ubangi")):
        v = S["t1"][pn]
        check("diapo %d cas" % sn, txt(sn, 11), n0(v["cases"]))
        check("diapo %d deces" % sn, txt(sn, 14), n0(v["deaths"]))
        check("diapo %d letalite" % sn, txt(sn, 16), "%.1f%%" % v["cfr"])
        check("diapo %d zones" % sn, txt(sn, 18), "%d / %d" % (v["hz_aff"], v["hz_tot"]))
    # alertes (tableau 3) : texte de la diapo 16
    t3 = S.get("t3_total")
    if t3:
        has("diapo 16 alertes", 16, 12, "%s alerts were received" % n0(t3["total"]))
        has("diapo 16 transferts", 16, 12, "%d were transferred" % t3["transferred"])
    # aucune valeur « NaN/None/%s » residuelle
    bad = []
    for i, sl in enumerate(p.slides, 1):
        for sh in sl.shapes:
            if sh.has_text_frame and re.search(r"\bNone\b|\bnan\b|%[sd]|\{|\}", sh.text_frame.text):
                bad.append("diapo %d forme %d : %r" % (i, sh.shape_id, sh.text_frame.text[:80]))
    if bad:
        fails.extend(["texte residuel suspect : " + b for b in bad])
    else:
        ok += 1
    import deck_tables
    import deck_nk
    nok, nfails = deck_nk.verify(p, S, db.rh)
    ok += nok
    fails.extend(nfails)
    tok, tfails = deck_tables.verify(p, S)
    ok += tok
    fails.extend(tfails)
    return ok, fails


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sitrep", type=int, required=True)
    ap.add_argument("--pdf-dir", default="data/pdf")
    ap.add_argument("--template", default="data/africa_cdc_deck/templates/deck_template.pptx")
    ap.add_argument("--state", default="data/africa_cdc_deck/state.json")
    ap.add_argument("--out-dir", default="outputs/deck_private")
    ap.add_argument("--narrative", default="auto", choices=["auto", "off"])
    ap.add_argument("--commit-state", action="store_true")
    ap.add_argument("--text", default=None, help="(tests) texte pdftotext deja extrait au lieu du PDF")
    ap.add_argument("--prev-text", default=None, help="(tests) texte du SitRep precedent")
    a = ap.parse_args()
    os.makedirs(a.out_dir, exist_ok=True)
    report = {"status": "error", "sitrep_no": a.sitrep}

    def finish(code, **kw):
        report.update(kw)
        json.dump(db.jsonable(report), open(os.path.join(a.out_dir, "deck_report.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print("[deck] %s" % json.dumps({k: report[k] for k in ("status", "sitrep_no") if k in report}), flush=True)
        return code

    try:
        if a.text:
            text = open(a.text, encoding="utf-8").read()
            pdf = a.text
        else:
            pdf = find_pdf(a.pdf_dir, a.sitrep)
            if not pdf:
                return finish(2, reason="PDF du SitRep %d introuvable dans %s" % (a.sitrep, a.pdf_dir))
            text = pdftotext(pdf)
        S = sp.parse_sitrep(text, a.sitrep)
    except sp.ParseError as e:
        return finish(2, status="refused", reason="SitRep %d illisible (format inattendu) : %s" % (a.sitrep, e))
    D = S["report_date"]
    state = load_json(a.state, {})
    C = state.get("config") or DEFAULT_CONFIG
    try:
        if a.prev_text:
            P = sp.parse_sitrep(open(a.prev_text, encoding="utf-8").read())
            pfp = a.prev_text
        else:
            P, pfp = find_prev(a.pdf_dir, D)
    except sp.ParseError as e:
        return finish(2, status="refused", reason="SitRep de la veille illisible : %s" % e)
    if not P:
        return finish(2, status="refused", reason="aucun SitRep du %s (veille) dans %s : comparaisons jour/jour impossibles, deck non genere" % ((D - dt.timedelta(days=1)).isoformat(), a.pdf_dir))
    P["isolation_by_prov"] = bool(P.get("care")) and all(P["care"].get(k, {}).get("isolation") is not None for k in P["care"]) and len(P["care"]) >= 4
    P["source_label"] = "SitRep %d" % P["sitrep_no"]
    H = load_history(a.pdf_dir) if not a.text else {}
    H[D.isoformat()] = (S["cases_cum"], S["deaths_cum"])
    narrative = {}
    if a.narrative == "auto" and os.environ.get("ANTHROPIC_API_KEY"):
        try:
            import narrative as nv
            narrative = nv.generate(S, text, state.get("narr") or {}, os.path.join(a.out_dir, "narrative_log.json"))
        except Exception as e:  # jamais bloquant
            print("[deck] narratif indisponible : %s" % e, flush=True)
            narrative = {}
    issue = S["sitrep_no"] - 10
    out_name = "BDBV_SitRep_%d_Dashboard_%d_%s_%d.pptx" % (issue, D.day, MOIS_FR[D.month], D.year)
    out_path = os.path.join(a.out_dir, out_name)
    try:
        dk, cx, wk = db.build(S, P, H, C, a.template, out_path, narrative=narrative, prior_narr=state.get("narr") or {})
    except db.DeckError as e:
        return finish(3 if "introuvable" in str(e) else 2, status="refused", reason=str(e))
    ok, fails = verify_deck(out_path, S, P, wk)
    ctl = os.path.join(a.out_dir, "Controle_%s.csv" % os.path.splitext(out_name)[0])
    with open(ctl, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Diapo", "Element", "Paragraphe/ligne", "Ancien", "Nouveau", "Remarque"])
        for r in dk.log:
            w.writerow(r)
    retained = [dict(slide=s_, element=e_, reason=r_) for s_, e_, r_ in dk.retained]
    narr_new = {k: v for k, v in (narrative or {}).items() if v}
    report.update(status="ok" if not fails else "verification_failed", file=out_name, control_csv=os.path.basename(ctl), issue=issue,
                  report_date=D, source_pdf=os.path.basename(pdf), prev_sitrep=P["sitrep_no"], edits=len(dk.log),
                  checks_passed=ok, checks_failed=fails, retained=retained, warnings=S["warnings"],
                  week=dict(epiweek=wk["epiweek"], last=wk.get("last"), so_far=wk.get("so_far")),
                  reconciliation=dict(cases="%d + %d = %d" % (P["cases_cum"], S["new_cases"], S["cases_cum"]),
                                      deaths="%d + %d = %d" % (P["deaths_cum"], S["new_deaths"], S["deaths_cum"]),
                                      cases_ok=(P["cases_cum"] + S["new_cases"] == S["cases_cum"]),
                                      deaths_ok=(P["deaths_cum"] + S["new_deaths"] == S["deaths_cum"])))
    if not (P["cases_cum"] + S["new_cases"] == S["cases_cum"]):
        report["warnings"] = list(report["warnings"]) + ["cumul de cas : veille %d + nouveaux %d != %d (revision ou donnee manquante)" % (P["cases_cum"], S["new_cases"], S["cases_cum"])]
    if not (P["deaths_cum"] + S["new_deaths"] == S["deaths_cum"]):
        report["warnings"] = list(report["warnings"]) + ["cumul de deces : veille %d + nouveaux %d != %d (revision ou donnee manquante)" % (P["deaths_cum"], S["new_deaths"], S["deaths_cum"])]
    code = 0 if not fails else 4
    if code == 0 and a.commit_state:
        st = dict(state)
        st["config"] = C
        st["last_built"] = dict(sitrep_no=S["sitrep_no"], report_date=D.isoformat(), file=out_name)
        narr = dict(state.get("narr") or {})
        narr.update(narr_new)
        st["narr"] = narr
        os.makedirs(os.path.dirname(a.state) or ".", exist_ok=True)
        json.dump(st, open(a.state, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return finish(code)


if __name__ == "__main__":
    sys.exit(main())
